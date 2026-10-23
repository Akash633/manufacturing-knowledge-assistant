import os
from pathlib import Path
from datetime import datetime, timezone
import docx
from docx.shared import Inches, Pt, RGBColor
import fitz # PyMuPDF
from PIL import Image, ImageDraw

WORKSPACE_ROOT = Path(__file__).resolve().parent.parent
CORPUS_DIR = WORKSPACE_ROOT / "data" / "synthetic_corpus"
CORPUS_DIR.mkdir(parents=True, exist_ok=True)

DISCLAIMER_BANNER = "SYNTHETIC DEMONSTRATION DATA. NOT VALID FOR REAL MACHINE OPERATION."

def create_sample_schematic_image(filepath: Path, label: str):
    """Creates a synthetic mechanical schematic image for multimodal testing."""
    img = Image.new("RGB", (600, 300), color=(240, 243, 246))
    draw = ImageDraw.Draw(img)
    
    # Draw border
    draw.rectangle([(10, 10), (590, 290)], outline=(30, 41, 59), width=3)
    # Draw axis diagram
    draw.line([(50, 250), (500, 250)], fill=(30, 41, 59), width=3)
    draw.line([(50, 250), (50, 50)], fill=(30, 41, 59), width=3)
    # Draw curve
    points = [(50, 230), (120, 210), (200, 150), (320, 110), (450, 90)]
    for i in range(len(points) - 1):
        draw.line([points[i], points[i+1]], fill=(14, 116, 144), width=4)

    draw.text((60, 60), f"DEMO-CNC-01 SCHEMATIC: {label}", fill=(15, 23, 42))
    draw.text((60, 80), DISCLAIMER_BANNER, fill=(185, 28, 28))
    img.save(filepath)

def generate_pdf_doc(
    filepath: Path,
    doc_number: str,
    title: str,
    revision: str,
    status: str,
    sections: list,
    schematic_path: Path = None,
):
    doc = fitz.open()

    for page_idx, (page_title, text_blocks) in enumerate(sections, start=1):
        page = doc.new_page(width=612, height=792) # Standard Letter size
        
        # Banner Header
        page.draw_rect(fitz.Rect(20, 20, 592, 45), color=(0.8, 0.1, 0.1), fill=(0.95, 0.9, 0.9))
        page.insert_text(fitz.Point(35, 36), DISCLAIMER_BANNER, fontsize=10, color=(0.7, 0.1, 0.1), fontname="helv")

        # Document Header
        header_text = f"DOCUMENT: {doc_number} | REV: {revision} | STATUS: {status.upper()}"
        page.insert_text(fitz.Point(35, 65), header_text, fontsize=11, color=(0.1, 0.2, 0.3), fontname="helv")
        page.insert_text(fitz.Point(35, 85), title, fontsize=14, color=(0.0, 0.2, 0.5), fontname="helv")
        page.draw_line(fitz.Point(35, 95), fitz.Point(577, 95), color=(0.6, 0.6, 0.6), width=1)

        y_offset = 120
        # Page Section Title
        page.insert_text(fitz.Point(35, y_offset), f"Page {page_idx}: {page_title}", fontsize=12, color=(0.1, 0.3, 0.5), fontname="helv")
        y_offset += 25

        for block in text_blocks:
            rect = fitz.Rect(35, y_offset, 577, y_offset + 120)
            page.insert_textbox(rect, block, fontsize=10, fontname="helv")
            y_offset += 110

        # Embed diagram on page 2 if provided
        if page_idx == 2 and schematic_path and schematic_path.exists():
            img_rect = fitz.Rect(100, y_offset, 500, y_offset + 180)
            page.insert_image(img_rect, filename=str(schematic_path))
            y_offset += 190
            page.insert_text(fitz.Point(120, y_offset + 15), f"Figure 1: {doc_number} Operational Limits Diagram", fontsize=9, color=(0.3, 0.3, 0.3), fontname="helv")

        # Footer
        page.draw_line(fitz.Point(35, 755), fitz.Point(577, 755), color=(0.7, 0.7, 0.7), width=1)
        footer_text = f"Fictional Machine: DEMO-CNC-01 | Page {page_idx} of {len(sections)} | Traceable Shop Floor System"
        page.insert_text(fitz.Point(35, 770), footer_text, fontsize=8, color=(0.4, 0.4, 0.4), fontname="helv")

    doc.save(str(filepath))
    doc.close()
    print(f"Generated PDF: {filepath.name}")

def generate_docx_doc(
    filepath: Path,
    doc_number: str,
    title: str,
    revision: str,
    status: str,
    sections: list,
    table_data: list = None,
):
    doc = docx.Document()
    
    # Disclaimer Banner
    p_disc = doc.add_paragraph()
    r_disc = p_disc.add_run(DISCLAIMER_BANNER)
    r_disc.bold = True
    r_disc.font.color.rgb = RGBColor(180, 20, 20)
    
    # Doc metadata
    h1 = doc.add_heading(f"{doc_number} ({revision}) - {status}", level=1)
    doc.add_paragraph(f"Title: {title}\nEffective Date: 2024-08-01\nMachine Model: DEMO-CNC-01")

    for sec_title, body_text in sections:
        doc.add_heading(sec_title, level=2)
        doc.add_paragraph(body_text)

    if table_data:
        doc.add_heading("Equipment Calibration & Specifications Table", level=2)
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        table.style = 'Table Grid'
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                cell = table.cell(r_idx, c_idx)
                cell.text = str(val)

    doc.save(str(filepath))
    print(f"Generated DOCX: {filepath.name}")

def main():
    schematic_path = CORPUS_DIR / "demo_spindle_schematic.png"
    create_sample_schematic_image(schematic_path, "Spindle Speed & Torque Envelope")

    # 1. Superseded SOP (Rev A)
    generate_pdf_doc(
        filepath=CORPUS_DIR / "SOP-CNC-001_RevA_SUPERSEDED.pdf",
        doc_number="SOP-CNC-001",
        title="Standard Operating Procedure: Spindle & Coolant Controls",
        revision="Rev A",
        status="Superseded",
        sections=[
            (
                "Section 1: General Machine Overview",
                [
                    "This document defines the basic startup and warmup routine for DEMO-CNC-01 3-axis machining center. All operators must hold Level 1 Machining Certification.",
                    "WARNING: This is Revision A, originally released 2021-03-01. Superseded by Revision B on 2024-01-15. Do not use for current operations."
                ]
            ),
            (
                "Section 2: Spindle Warm-up Routine (Legacy)",
                [
                    "Spindle Warm-up Procedure (Rev A Legacy):\nExecute a single 15-minute warmup sequence at constant 1,000 RPM. Verify spindle chiller oil temperature reads between 20°C and 25°C.",
                    "Legacy Coolant Parameters: Maintain water-soluble synthetic coolant at 5% to 7% Brix concentration. Flood coolant pressure set to 1.5 bar."
                ]
            ),
            (
                "Section 3: Legacy Cutting Parameters",
                [
                    "Feed Rate Guidelines:\nFor general roughing on 6061-T6 Aluminum, feed rate is 2,000 mm/min at 8,000 RPM.\nFor Titanium Grade 5 (Ti-6Al-4V), maximum recommended feed rate is 1,500 mm/min.",
                    "Safety Note: Ensure door interlock switches are engaged before cycle start."
                ]
            ),
        ],
        schematic_path=schematic_path,
    )

    # 2. Approved SOP (Rev B)
    generate_pdf_doc(
        filepath=CORPUS_DIR / "SOP-CNC-001_RevB_APPROVED.pdf",
        doc_number="SOP-CNC-001",
        title="Standard Operating Procedure: Spindle & Coolant Controls",
        revision="Rev B",
        status="Approved",
        sections=[
            (
                "Section 1: Scope, Approval & Revision Summary",
                [
                    "Official Standard Operating Procedure for DEMO-CNC-01 Vertical Machining Center. Approved by Quality Engineering and Operations Director on 2024-01-15.",
                    "Revision B replaces and supersedes Rev A in its entirety. Spindle warmup cycle updated to dual-stage sequence to eliminate ceramic bearing thermal shock."
                ]
            ),
            (
                "Section 2: Mandatory Spindle Warm-up Cycle (Approved)",
                [
                    "Spindle Warm-up Procedure (Approved Rev B):\nStage 1: Run spindle at 1,500 RPM for exactly 5 minutes with flood coolant circulating.\nStage 2: Ramp spindle to 3,000 RPM for exactly 5 minutes.\nTotal warm-up duration: 10 minutes.",
                    "Coolant Concentration: Check coolant tank using optical refractometer every shift. Maintain 8.0% to 10.0% concentration for tough alloys (Titanium/Inconel). Minimum flood coolant pump pressure: 2.5 bar."
                ]
            ),
            (
                "Section 3: Catalog Feed & Speed References",
                [
                    "Catalog Machining Parameters:\nAluminum 6061-T6: 10,000 RPM, feed rate 2,400 mm/min.\nTitanium Grade 5 (Ti-6Al-4V) general milling catalog value: 3,500 RPM, feed rate 1,500 mm/min.\n(Note: Always cross-reference active Quality Alerts before initiating cuts on exotic alloys).",
                    "Critical Safety Mandate: Enclosure doors must remain closed and interlocked during spindle rotation. Never bypass safety interlocks. Lockout/Tagout (LOTO) is mandatory before opening electrical cabinets."
                ]
            ),
        ],
        schematic_path=schematic_path,
    )

    # 3. Preventive Maintenance Report (DOCX)
    generate_docx_doc(
        filepath=CORPUS_DIR / "MR-2024-08_RevA_APPROVED.docx",
        doc_number="MR-2024-08",
        title="Preventive Maintenance & Calibration Report",
        revision="Rev A",
        status="Approved",
        sections=[
            (
                "Section 1: Axis Backlash & Laser Interferometer Verification",
                "Periodic 6-month calibration performed on DEMO-CNC-01 by Lead Maintenance Specialist. "
                "Ball screws inspected for pitch error and backlash. Maximum allowable X-axis backlash tolerance is 0.008 mm. "
                "Actual measured backlash on X-axis: 0.004 mm (PASS). Y-axis backlash: 0.003 mm (PASS). Z-axis backlash: 0.005 mm (PASS)."
            ),
            (
                "Section 2: Lubrication Specifications & Way Lube",
                "Slideway lubrication system requires Mobil Vactra No. 2 (ISO VG 68 way oil). "
                "Inspect reservoir fluid level every morning before shift start. Pneumatic pressure regulator set to 6.0 bar. "
                "Spindle vibration analysis registered 0.9 mm/s RMS, well below ISO 10816-1 threshold of 1.8 mm/s."
            ),
        ],
        table_data=[
            ["Machine Parameter", "Specification Limit", "Measured Calibration", "Status"],
            ["X-Axis Backlash", "Max 0.008 mm", "0.004 mm", "PASSED"],
            ["Y-Axis Backlash", "Max 0.008 mm", "0.003 mm", "PASSED"],
            ["Z-Axis Backlash", "Max 0.010 mm", "0.005 mm", "PASSED"],
            ["Spindle Vibration RMS", "Max 1.8 mm/s", "0.9 mm/s", "PASSED"],
            ["Way Lube Type", "ISO VG 68 (Vactra 2)", "Verified", "PASSED"],
        ]
    )

    # 4. Quality Alert (PDF)
    generate_pdf_doc(
        filepath=CORPUS_DIR / "QA-CNC-2024-03_RevA_APPROVED.pdf",
        doc_number="QA-CNC-2024-03",
        title="Quality Alert: Titanium Grade 5 Chatter & Feed Rate Restriction",
        revision="Rev A",
        status="Approved",
        sections=[
            (
                "Section 1: Problem Description & Defect Root Cause",
                [
                    "QUALITY ALERT QA-CNC-2024-03 | Effective Date: 2024-03-10 | Severity: HIGH PRIORITY",
                    "Defect Summary: Surface chatter marks exceeding Ra 1.6 μm and micro-chipping on Ti-6Al-4V turbine bracket flanges produced on DEMO-CNC-01. Root cause analysis determined that high cutting forces at 1,500 mm/min feed rate induced harmonic resonance in workpiece fixtures."
                ]
            ),
            (
                "Section 2: Mandatory Operational Restriction",
                [
                    "MANDATORY OVERRIDE FOR ALL DEMO-CNC-01 OPERATORS:\n"
                    "Maximum allowable feed rate for face milling Titanium Grade 5 is strictly capped at 1,200 mm/min (Reduced from general catalog limit of 1,500 mm/min).\n"
                    "Do NOT exceed 1,200 mm/min feed rate under any circumstance until fixture redesign (ECN-892) is released.",
                    "Tooling Requirement: Use 4-flute solid carbide endmill with 0.5 mm corner radius and minimum 8% water-soluble coolant flow."
                ]
            ),
        ],
        schematic_path=schematic_path,
    )

    print("All synthetic demonstration files generated successfully.")

if __name__ == "__main__":
    main()
